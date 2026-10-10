"""End-to-end tests against the running docker-compose stack (BASE_URL).

Every conversion is checked for content (text in logical order, Arabic
included) and for looks (rendered pages compared pixel-wise with the source).
"""
import difflib
import io
import zipfile

import numpy as np
import pymupdf
import pytest
from docx import Document
from PIL import Image

from conftest import FIX, pdf_text, render, save_side_by_side, squash, visual_match

DOCX_MT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def post(api, path, files, **data):
    r = api.post(path, files=files, data={k: str(v) for k, v in data.items()})
    assert r.status_code == 200, f"{path}: {r.status_code} {r.text[:500]}"
    return r


def ocr_readback(api, pdf: bytes, dpi: int = 200) -> str:
    """Rasterise a PDF and OCR it again: proves the *rendered* Arabic glyphs are
    correctly shaped and ordered (scrambled / disconnected Arabic does not OCR)."""
    img = Image.fromarray(render(pdf, dpi))
    buf = io.BytesIO()
    img.save(buf, "PNG", dpi=(dpi, dpi))
    return post(api, "/api/v1/ocr/text", {"file": ("page.png", buf.getvalue())}).text


def docx_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    # invisible left-to-right marks keep Latin / numbers in their visual order
    return "\n".join(parts).replace("\u200e", "")


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_structure(data: bytes) -> dict:
    """Native content of a .docx: shapes, anchored pictures (position / size in pt),
    tables, tab stops, and Word schema-order problems."""
    import re
    import zipfile

    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    pics = []
    for m in re.finditer(r"<wp:anchor.*?</wp:anchor>", xml, re.S):
        a = m.group(0)
        if "<pic:pic" not in a:
            continue
        x, y = (int(v) / 12700 for v in re.findall(r"<wp:posOffset>(-?\d+)</wp:posOffset>", a)[:2])
        cx, cy = (int(v) / 12700 for v in re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"', a).groups())
        pics.append((x, y, cx, cy))
    problems = []
    d = Document(io.BytesIO(data))
    order = {"tcPr": ["tcW", "gridSpan", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "vAlign"],
             "tblPr": ["tblStyle", "bidiVisual", "tblW", "jc", "tblInd", "tblBorders", "tblLayout",
                       "tblCellMar", "tblLook"]}
    for name, seq in order.items():
        for el in d.element.body.iter(W_NS + name):
            tags = [c.tag.replace(W_NS, "") for c in el]
            if len(tags) != len(set(tags)):
                problems.append(f"duplicate in {name}: {tags}")
            known = [t for t in tags if t in seq]
            if known != sorted(known, key=seq.index):
                problems.append(f"order in {name}: {tags}")
    return {"shapes": xml.count("<wps:wsp>"), "pictures": pics, "tables": xml.count("<w:tbl>"),
            "tabs": xml.count("<w:tab "), "problems": problems}


def docx_looks_like(api, docx: bytes, reference: np.ndarray, dpi: float, out_png=None) -> dict:
    """Render the .docx with the service (LibreOffice in Gotenberg) and compare it
    pixel-wise with the reference page."""
    pdf = post(api, "/api/v1/word-to-pdf", {"file": ("x.docx", docx)}).content
    assert pymupdf.open(stream=pdf, filetype="pdf").page_count == 1
    got = render(pdf, dpi)
    if out_png:
        save_side_by_side(reference, got, out_png)
    return visual_match(reference, got, tol=2)


def docx_baseline_error(api, docx: bytes, reference_pdf: bytes) -> tuple[float, float]:
    """Render the .docx back to PDF and compare every line baseline with the source."""
    pdf = post(api, "/api/v1/word-to-pdf", {"file": ("x.docx", docx)}).content

    def lines(b):
        out = []
        for blk in pymupdf.open(stream=b, filetype="pdf")[0].get_text("dict")["blocks"]:
            for ln in blk.get("lines", []):
                if "".join(s["text"] for s in ln["spans"]).strip():
                    out.append((ln["spans"][0]["origin"][1], ln["bbox"][0]))
        return out

    ref, got = lines(reference_pdf), lines(pdf)
    errs = [min(abs(g[0] - r[0]) for g in got) for r in ref]
    return max(errs), float(np.mean(errs))


# ===================================================================== meta
def test_health(api):
    j = api.get("/health").json()
    assert j["status"] == "up"
    assert j["gotenberg"]["details"]["chromium"]["status"] == "up"
    assert j["gotenberg"]["details"]["libreoffice"]["status"] == "up"


def test_bad_input_is_rejected(api):
    r = api.post("/api/v1/ocr/pdf", files={"file": ("x.png", b"not an image")})
    assert r.status_code == 400
    r = api.post("/api/v1/pdf-to-image", files={"file": ("x.pdf", b"%PDF-broken")})
    assert r.status_code == 400
    r = api.post("/api/v1/word-to-pdf", files={"file": ("x.exe", b"MZ")})
    assert r.status_code == 400


# ===================================================================== word -> pdf
def test_word_to_pdf_arabic(api, sample_pdf, truth, out_dir):
    doc = pymupdf.open(stream=sample_pdf, filetype="pdf")
    assert doc.page_count == 1
    fonts = {f[3] for f in doc[0].get_fonts()}
    assert any("Naskh" in f for f in fonts), fonts          # Arabic font embedded
    text = squash(pdf_text(sample_pdf))
    for phrase in truth["sample_ar.docx"]:
        assert squash(phrase) in text, phrase                # logical order text layer
    # the rendered page must OCR back to the same Arabic: shaping + RTL are right
    back = squash(ocr_readback(api, sample_pdf))
    for phrase in ["تقرير الخدمات الرقمية", "تهدف هذه الوثيقة", "من اليمين إلى",
                   "داخل النص", "This paragraph is written in English"]:
        assert squash(phrase) in back, phrase
    Image.fromarray(render(sample_pdf, 100)).save(out_dir / "word_to_pdf.png")


# ===================================================================== pdf -> word
def test_pdf_to_word_arabic(api, sample_pdf, truth, out_dir):
    r = post(api, "/api/v1/pdf-to-word", {"file": ("sample_ar.pdf", sample_pdf)})
    assert r.headers["content-type"] == DOCX_MT
    (out_dir / "pdf_to_word.docx").write_bytes(r.content)
    d = Document(io.BytesIO(r.content))
    text = docx_text(r.content)
    # exact logical text, punctuation included
    assert "تقرير الخدمات الرقمية" in text
    assert "يحتوي هذا السطر على أرقام 2024 وكلمات إنجليزية مثل PDF و Word داخل النص العربي." in text
    assert "This paragraph is written in English to verify mixed-direction documents." in text
    # right-to-left paragraphs and runs
    xml = d.element.xml
    assert "<w:bidi/>" in xml and "<w:rtl/>" in xml
    # the table is a real Word table, right-to-left, 3 x 3
    assert len(d.tables) == 1
    t = d.tables[0]
    assert len(t.rows) == 3 and len(t.columns) == 3
    assert "<w:bidiVisual/>" in t._tbl.xml
    assert [c.text for c in t.rows[0].cells] == ["الاسم", "الكمية", "السعر"]
    assert [c.text for c in t.rows[2].cells] == ["دفتر", "2", "40"]
    # layout: every line baseline lands where it was in the PDF, and it looks the same
    worst, mean = docx_baseline_error(api, r.content, sample_pdf)
    assert worst < 1.0 and mean < 0.5, (worst, mean)
    st = docx_structure(r.content)
    assert not st["problems"] and st["pictures"] == [] and st["shapes"] == 0, st
    m = docx_looks_like(api, r.content, render(sample_pdf, 150), 150, out_dir / "pdf_to_word_compare.png")
    assert m["precision"] > 0.98 and m["recall"] > 0.98, m


# ===================================================================== image -> pdf
def test_image_to_pdf_lossless_jpeg(api):
    jpg = (FIX / "photo.jpg").read_bytes()
    pdf = post(api, "/api/v1/image-to-pdf", {"files": ("photo.jpg", jpg)}).content
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert doc.page_count == 1
    xref = doc[0].get_images()[0][0]
    assert doc.xref_stream_raw(xref) == jpg               # JPEG embedded bit-exact


def test_image_to_pdf_png_pixels_and_size(api):
    png = (FIX / "invoice_ar.png").read_bytes()
    pdf = post(api, "/api/v1/image-to-pdf", {"files": ("invoice.png", png)}).content
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    page = doc[0]
    src = Image.open(io.BytesIO(png))
    dpi = src.info["dpi"][0]
    assert abs(page.rect.width - src.width * 72 / dpi) < 1      # physical size kept
    assert abs(page.rect.height - src.height * 72 / dpi) < 1
    info = doc.extract_image(page.get_images()[0][0])
    emb = np.asarray(Image.open(io.BytesIO(info["image"])).convert("RGB"))
    assert np.array_equal(emb, np.asarray(src.convert("RGB")))  # pixel exact


def test_images_to_pdf_multi_a4(api):
    files = [("files", ("a.jpg", (FIX / "photo.jpg").read_bytes())),
             ("files", ("b.png", (FIX / "letter_ar.png").read_bytes()))]
    pdf = post(api, "/api/v1/image-to-pdf", files, page_size="A4", margin=20).content
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert doc.page_count == 2
    # landscape photo -> landscape A4 (auto orientation), portrait letter -> portrait
    assert round(doc[0].rect.width) == 842 and round(doc[0].rect.height) == 595
    assert round(doc[1].rect.width) == 595 and round(doc[1].rect.height) == 842


# ===================================================================== pdf -> image
def test_pdf_to_image_single_and_zip(api, sample_pdf):
    r = post(api, "/api/v1/pdf-to-image", {"file": ("s.pdf", sample_pdf)}, format="png", dpi=144)
    assert r.headers["content-type"] == "image/png"
    img = Image.open(io.BytesIO(r.content))
    page = pymupdf.open(stream=sample_pdf, filetype="pdf")[0]
    assert abs(img.width - page.rect.width * 2) <= 1
    ref = render(sample_pdf, 144)
    assert np.array_equal(np.asarray(img.convert("RGB")), ref)

    two = pymupdf.open()
    two.insert_pdf(pymupdf.open(stream=sample_pdf, filetype="pdf"))
    two.insert_pdf(pymupdf.open(stream=sample_pdf, filetype="pdf"))
    r = post(api, "/api/v1/pdf-to-image", {"file": ("two.pdf", two.tobytes())}, format="jpg", dpi=72)
    assert r.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert names == ["page-001.jpg", "page-002.jpg"]


# ===================================================================== OCR -> regenerated PDF
@pytest.mark.parametrize("name", ["invoice_ar.png", "letter_ar.png"])
def test_ocr_regenerated_pdf_looks_like_original(api, truth, out_dir, name):
    src = (FIX / name).read_bytes()
    r = post(api, "/api/v1/ocr/pdf", {"file": (name, src)})
    pdf = r.content
    (out_dir / f"ocr_{name}.pdf").write_bytes(pdf)
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    page = doc[0]
    orig = np.asarray(Image.open(io.BytesIO(src)).convert("RGB"))
    dpi = Image.open(io.BytesIO(src)).info["dpi"][0]

    # 1. same physical page size as the scan
    assert abs(page.rect.width - orig.shape[1] * 72 / dpi) < 1
    # 2. regenerated, not a picture: real embedded fonts, and no image covers the page
    assert any("Naskh" in f[3] for f in page.get_fonts())
    for info in page.get_image_info():
        x0, y0, x1, y1 = info["bbox"]
        assert (x1 - x0) * (y1 - y0) < 0.2 * page.rect.width * page.rect.height
    # 3. text layer: every expected phrase, in logical order
    text = squash(pdf_text(pdf))
    for phrase in truth[name]:
        assert squash(phrase) in text, phrase
    # 4. looks the same: ink overlap and colour error against the original scan
    regen = render(pdf, dpi)
    m = visual_match(orig, regen)
    save_side_by_side(orig, regen, out_dir / f"ocr_{name}_compare.png")
    assert m["precision"] > 0.93 and m["recall"] > 0.93, m
    assert m["mae"] < 6.0, m
    # 5. the regenerated glyphs OCR back to the same text as the original scan does
    #    (broken shaping / reversed Arabic would collapse this similarity)
    first = squash(post(api, "/api/v1/ocr/text", {"file": (name, src)}).text)
    back = squash(ocr_readback(api, pdf, round(dpi)))
    ratio = difflib.SequenceMatcher(None, first, back).ratio()
    assert ratio > 0.97, (ratio, first, back)


def test_ocr_pdf_from_scanned_pdf_and_background_modes(api, out_dir):
    # scanned pdf input (image inside a pdf)
    scan = post(api, "/api/v1/image-to-pdf", {"files": ("inv.png", (FIX / "invoice_ar.png").read_bytes())}).content
    for bg in ("vector", "original", "none"):
        pdf = post(api, "/api/v1/ocr/pdf", {"file": ("scan.pdf", scan)}, background=bg).content
        (out_dir / f"ocr_scan_{bg}.pdf").write_bytes(pdf)
        text = squash(pdf_text(pdf))
        assert squash("المجموع الكلي") in text and "19750" in text


# ===================================================================== OCR -> word / text / json
def test_ocr_to_word(api, truth, out_dir):
    src = (FIX / "invoice_ar.png").read_bytes()
    r = post(api, "/api/v1/ocr/word", {"file": ("invoice_ar.png", src)})
    (out_dir / "ocr_invoice.docx").write_bytes(r.content)
    d = Document(io.BytesIO(r.content))
    text = docx_text(r.content)
    for phrase in ["فاتورة ضريبية", "اسم العميل: شركة النور للتجارة", "العنوان: الرباط، المغرب",
                   "تاريخ الاستحقاق 15 مارس 2025", "المجموع الكلي: 19750 درهم",
                   "Thank you for your business. Invoice No. 2025-0042"]:
        assert phrase in text, phrase
    assert len(d.tables) == 1
    t = d.tables[0]
    assert len(t.rows) == 4 and len(t.columns) == 3
    assert [c.text for c in t.rows[0].cells] == ["المنتج", "الكمية", "السعر"]
    assert [c.text for c in t.rows[1].cells] == ["حاسوب محمول", "2", "8500"]
    assert "<w:bidiVisual/>" in t._tbl.xml
    # header cells keep their shading
    assert 'w:fill="CF' in t._tbl.xml or 'w:fill="D0' in t._tbl.xml
    st = docx_structure(r.content)
    assert not st["problems"], st["problems"]
    # band as a native shape, the logo as the only picture, exactly where it was
    assert st["shapes"] >= 1
    assert len(st["pictures"]) == 1, st["pictures"]
    x, y, w, h = st["pictures"][0]
    k = 72 / 200
    lx0, ly0, lx1, ly1 = 99 * k, 58 * k, 312 * k, 271 * k     # logo ink box in the scan
    assert x <= lx0 + 1 and y <= ly0 + 1 and x + w >= lx1 - 1 and y + h >= ly1 - 1
    assert w < (lx1 - lx0) + 6 and h < (ly1 - ly0) + 6
    # rendered back, it looks like the scan
    orig = np.asarray(Image.open(io.BytesIO(src)).convert("RGB"))
    m = docx_looks_like(api, r.content, orig, 200, out_dir / "ocr_word_compare.png")
    assert m["precision"] > 0.98 and m["recall"] > 0.98, m


def test_ocr_text_and_json(api, truth):
    src = (FIX / "letter_ar.png").read_bytes()
    text = post(api, "/api/v1/ocr/text", {"file": ("l.png", src)}).text
    assert "تقرير الخدمات الرقمية" in text
    assert "يحتوي هذا السطر على أرقام 2024 وكلمات إنجليزية مثل PDF و Word داخل النص" in text
    j = post(api, "/api/v1/ocr/json", {"file": ("l.png", src)}).json()
    lines = j["pages"][0]["lines"]
    title = lines[0]
    assert title["rtl"] and title["text"] == "تقرير الخدمات الرقمية"
    assert all(w["bold"] for w in title["words"])
    assert title["words"][0]["color"].startswith("#")


def test_pdf_to_word_scanned_pdf_uses_ocr(api):
    scan = post(api, "/api/v1/image-to-pdf", {"files": ("l.png", (FIX / "letter_ar.png").read_bytes())}).content
    r = post(api, "/api/v1/pdf-to-word", {"file": ("scan.pdf", scan)})
    text = docx_text(r.content)
    assert "تقرير الخدمات الرقمية" in text
    assert "This paragraph is written in English to verify mixed-direction documents." in text


# ===================================================================== hard input
def test_ocr_phone_capture(api, out_dir):
    """Tilted, noisy, JPEG q70, ~160 dpi, no DPI metadata."""
    src = (FIX / "phone_invoice.jpg").read_bytes()
    for deskew in (False, True):
        pdf = post(api, "/api/v1/ocr/pdf", {"file": ("phone.jpg", src)}, deskew=deskew).content
        (out_dir / f"ocr_phone_deskew_{deskew}.pdf").write_bytes(pdf)
        text = squash(pdf_text(pdf))
        for phrase in ["فاتورة ضريبية", "اسم العميل", "شركة النور للتجارة", "الرباط", "المغرب",
                       "حاسوب محمول", "طابعة ليزر", "المجموع الكلي", "19750", "8500", "2300",
                       "Thank you for your business", "2025-0042"]:
            assert squash(phrase) in text, (deskew, phrase)
        page = pymupdf.open(stream=pdf, filetype="pdf")[0]
        # the round logo must stay a picture, never become a giant letter
        sizes = [s["size"] for b in page.get_text("dict")["blocks"] for ln in b.get("lines", [])
                 for s in ln["spans"]]
        assert max(sizes) < 40, max(sizes)
        if not deskew:   # without deskew the output must look exactly like the photo
            orig = np.asarray(Image.open(io.BytesIO(src)).convert("RGB"))
            dpi = orig.shape[1] / (page.rect.width / 72)
            regen = render(pdf, dpi)
            save_side_by_side(orig, regen, out_dir / "ocr_phone_compare.png")
            m = visual_match(orig, regen, tol=3)
            assert m["precision"] > 0.9 and m["recall"] > 0.9, m


# ===================================================================== native Word graphics
def test_pdf_to_word_graphics_are_native(api, out_dir):
    """Digital PDF with a vector band, a vector logo, a table, a photo and a drawn
    signature: band -> native shape, table -> native table, logo / photo / signature
    -> pictures at their exact place, all text editable."""
    src = (FIX / "invoice_ar.pdf").read_bytes()
    r = post(api, "/api/v1/pdf-to-word", {"file": ("invoice_ar.pdf", src)})
    (out_dir / "pdf_to_word_invoice.docx").write_bytes(r.content)
    st = docx_structure(r.content)
    assert not st["problems"], st["problems"]
    assert st["tables"] == 1 and st["shapes"] >= 1
    assert len(st["pictures"]) == 3, st["pictures"]
    page = pymupdf.open(stream=src, filetype="pdf")[0]
    photo = pymupdf.Rect(page.get_image_info()[0]["bbox"])
    assert any(abs(x - photo.x0) < 1 and abs(y - photo.y0) < 1 and abs(w - photo.width) < 1
               and abs(h - photo.height) < 1 for x, y, w, h in st["pictures"]), st["pictures"]
    # nothing is a page-sized picture
    assert all(w * h < 0.2 * page.rect.width * page.rect.height for _, _, w, h in st["pictures"])
    text = docx_text(r.content)
    for phrase in ["فاتورة ضريبية", "اسم العميل: شركة النور للتجارة", "المجموع الكلي: 19750 درهم",
                   "التوقيع", "Thank you for your business. Invoice No. 2025-0042"]:
        assert phrase in text, phrase
    m = docx_looks_like(api, r.content, render(src, 150), 150, out_dir / "pdf_to_word_invoice_compare.png")
    assert m["precision"] > 0.98 and m["recall"] > 0.98, m


def test_pdf_to_word_form_uses_tab_stops(api, out_dir):
    """Label / value pairs spread over a line and an unruled table keep their exact
    horizontal positions with native tab stops (Arabic numbers keep their order)."""
    src = (FIX / "form_ar.pdf").read_bytes()
    r = post(api, "/api/v1/pdf-to-word", {"file": ("form_ar.pdf", src)})
    (out_dir / "pdf_to_word_form.docx").write_bytes(r.content)
    st = docx_structure(r.content)
    assert not st["problems"] and st["tabs"] >= 6 and st["pictures"] == [], st
    text = docx_text(r.content)
    for phrase in ["الاسم: أحمد بن علي", "رقم الملف:", "Reference: AB-77", "استشارة قانونية", "1200"]:
        assert phrase in text, phrase
    m = docx_looks_like(api, r.content, render(src, 150), 150, out_dir / "pdf_to_word_form_compare.png")
    assert m["precision"] > 0.98 and m["recall"] > 0.98, m


def test_ocr_to_word_phone_capture(api, out_dir):
    """Tilted noisy photo: straightened automatically for Word, table rebuilt natively."""
    src = (FIX / "phone_invoice.jpg").read_bytes()
    r = post(api, "/api/v1/ocr/word", {"file": ("phone.jpg", src)})
    (out_dir / "ocr_phone.docx").write_bytes(r.content)
    st = docx_structure(r.content)
    assert not st["problems"] and st["tables"] == 1, st
    d = Document(io.BytesIO(r.content))
    t = d.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["المنتج", "الكمية", "السعر"]
    text = docx_text(r.content)
    assert "المجموع الكلي: 19750 درهم" in text
